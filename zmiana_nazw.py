#!/usr/bin/env python3
"""Masowa zmiana nazw plikow w folderze: wedlug wzoru ({nazwa}, {nr}, {data},
{folder}, zamiana tekstu) albo wedlug listy w Excelu (kolumna A -> B).
Podglad przed zmiana, wykrywanie konfliktow, cofniecie ostatniej zmiany.

Uruchomienie: python zmiana_nazw.py            (GUI)
              python zmiana_nazw.py --selftest (test logiki)
"""
import csv
import datetime
import fnmatch
import glob
import os
import re
import sys
import traceback
import uuid

if getattr(sys, "frozen", False):
    APP_DIR = os.path.dirname(sys.executable)
else:
    APP_DIR = os.path.dirname(os.path.abspath(__file__))
HIST_DIR = os.path.join(APP_DIR, "OUTPUT")

BAD_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
RESERVED = {"CON", "PRN", "AUX", "NUL", *("COM%d" % i for i in range(1, 10)),
            *("LPT%d" % i for i in range(1, 10))}
FIELD = re.compile(r"\{(\w+)(?::(\d+))?\}")
ORDERS = ("nazwa", "data modyfikacji")

# ---------------------------------------------------------------- pliki ----


def natural_key(s):
    """'skan2' przed 'skan10'."""
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", s)]


def list_files(folder, mask="*", order="nazwa"):
    names = [f for f in os.listdir(folder)
             if os.path.isfile(os.path.join(folder, f))
             and fnmatch.fnmatch(f.lower(), (mask.strip() or "*").lower())]
    if order == "data modyfikacji":
        return sorted(names, key=lambda f: (os.path.getmtime(os.path.join(folder, f)),
                                            natural_key(f)))
    return sorted(names, key=natural_key)


def name_error(name):
    """Powod, dla ktorego Windows nie przyjmie nazwy, albo None."""
    if not name.strip() or name.startswith("."):
        return "pusta nazwa"
    if BAD_CHARS.search(name):
        return 'niedozwolony znak (\\ / : * ? " < > |)'
    if name != name.rstrip(" ."):
        return "nazwa nie może kończyć się spacją ani kropką"
    if name.split(".")[0].upper().strip() in RESERVED:
        return "nazwa zarezerwowana przez Windows"
    return None


# ----------------------------------------------------------------- plany ----


def plan_pattern(folder, names, pattern="{nazwa}", find="", repl="", start=1):
    """[(stara, nowa)]. Rozszerzenie zostaje. Pola wzoru: {nazwa} {nr} {nr:3}
    {data} (data modyfikacji RRRR-MM-DD) {folder}."""
    pattern = pattern.strip() or "{nazwa}"
    for m in FIELD.finditer(pattern):
        if m.group(1) not in ("nazwa", "nr", "data", "folder"):
            raise ValueError("Nieznane pole {%s} - dostepne: {nazwa} {nr} {nr:3} {data} {folder}"
                             % m.group(1))
    plan = []
    for i, old in enumerate(names):
        stem, ext = os.path.splitext(old)
        if find:
            stem = stem.replace(find, repl)
        mtime = os.path.getmtime(os.path.join(folder, old))
        values = {"nazwa": stem,
                  "data": datetime.date.fromtimestamp(mtime).isoformat(),
                  "folder": os.path.basename(os.path.abspath(folder))}

        def field(m):
            if m.group(1) == "nr":
                return str(start + i).zfill(int(m.group(2) or 1))
            return values[m.group(1)]
        plan.append((old, FIELD.sub(field, pattern).strip() + ext))
    return plan


def write_list(folder, names, xlsx):
    """Lista plikow do Excela: A = obecna nazwa, B = nowa nazwa (do wpisania)."""
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Zmiana nazw"
    ws.append(["Obecna nazwa", "Nowa nazwa"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for n in names:
        ws.append([n, None])
    ws.column_dimensions["A"].width = ws.column_dimensions["B"].width = 50
    wb.save(xlsx)


def plan_excel(xlsx):
    """[(stara, nowa)] z kolumn A i B (wiersz 1 = naglowek, puste B pomijane).
    Brak rozszerzenia w nowej nazwie -> dopisywane rozszerzenie starej."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    try:
        rows = list(wb.worksheets[0].iter_rows(min_row=2, max_col=2, values_only=True))
    finally:
        wb.close()
    plan = []
    for row in rows:
        old, new = (str(v).strip() if v is not None else "" for v in (row + (None, None))[:2])
        if not old or not new:
            continue
        ext = os.path.splitext(old)[1]
        if ext and not new.lower().endswith(ext.lower()):
            new += ext
        plan.append((old, new))
    return plan


def check(folder, plan):
    """[(stara, nowa, uwaga)]. Uwaga '' = OK, 'bez zmian' = pomijane,
    inna = blad (nic nie zostanie zmienione, dopoki bledy sa)."""
    olds = {o.lower() for o, _ in plan}
    targets = {}
    for _, n in plan:
        targets[n.lower()] = targets.get(n.lower(), 0) + 1
    out = []
    for old, new in plan:
        if not os.path.isfile(os.path.join(folder, old)):
            note = "brak takiego pliku w folderze"
        elif old == new:
            note = "bez zmian"
        else:
            note = name_error(new) or ""
            if not note and targets[new.lower()] > 1:
                note = "dwa pliki dostałyby tę samą nazwę"
            elif (not note and new.lower() not in olds
                  and os.path.exists(os.path.join(folder, new))):
                note = "plik o tej nazwie już istnieje"
        out.append((old, new, note))
    return out


def errors(checked):
    return [c for c in checked if c[2] not in ("", "bez zmian")]


# ----------------------------------------------------------- zmiana nazw ----


def _rename_all(folder, pairs):
    """Zmiana nazw dwufazowa (przez nazwy tymczasowe), wiec zamiana miejscami
    a<->b i lancuchy a->b->c dzialaja. Blad = cofniecie juz zrobionych."""
    tag = ".~zn" + uuid.uuid4().hex[:8]
    steps = [(o, "%s%s%d" % (o, tag, i)) for i, (o, _) in enumerate(pairs)]
    steps += [(t, n) for (_, t), (_, n) in zip(steps, pairs)]
    done = []
    try:
        for a, b in steps:
            os.rename(os.path.join(folder, a), os.path.join(folder, b))
            done.append((a, b))
    except OSError:
        for a, b in reversed(done):
            try:
                os.rename(os.path.join(folder, b), os.path.join(folder, a))
            except OSError:
                pass
        raise


def apply(folder, plan, hist_dir=HIST_DIR):
    """Zmienia nazwy (tylko gdy brak bledow). Zapisuje historie do cofniecia.
    Zwraca liczbe zmienionych plikow."""
    checked = check(folder, plan)
    if errors(checked):
        raise ValueError("W planie sa bledy - popraw je przed zmiana nazw.")
    pairs = [(o, n) for o, n, note in checked if note == ""]
    if not pairs:
        return 0
    _rename_all(folder, pairs)
    os.makedirs(hist_dir, exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    with open(os.path.join(hist_dir, "historia_%s.csv" % stamp), "w",
              newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["folder", "stara nazwa", "nowa nazwa"])
        w.writerows((os.path.abspath(folder), o, n) for o, n in pairs)
    return len(pairs)


def last_change(hist_dir=HIST_DIR):
    """(plik_historii, folder, [(nowa, stara)]) ostatniej zmiany albo None."""
    files = sorted(glob.glob(os.path.join(hist_dir, "historia_*.csv")))
    if not files:
        return None
    with open(files[-1], newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f, delimiter=";"))[1:]
    return files[-1], rows[0][0], [(n, o) for _, o, n in rows]


def undo_last(hist_dir=HIST_DIR):
    """Cofa ostatnia zmiane z historii. Zwraca (folder, liczba) albo None."""
    last = last_change(hist_dir)
    if not last:
        return None
    path, folder, back = last
    bad = errors(check(folder, back))
    if bad:
        raise ValueError("Nie mozna cofnac - pliki zmienily sie od tego czasu: %s (%s)"
                         % (bad[0][0], bad[0][2]))
    _rename_all(folder, back)
    os.replace(path, os.path.join(hist_dir, "cofnieta_" + os.path.basename(path)))
    return folder, len(back)


# -------------------------------------------------------------------- GUI ----


def gui():
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    root = tk.Tk()
    root.title("Zmiana nazw plikow")
    root.geometry("900x640")
    pad = dict(padx=6, pady=3)

    v_dir = tk.StringVar(value=os.path.join(APP_DIR, "INPUT"))
    v_mask = tk.StringVar(value="*")
    v_order = tk.StringVar(value=ORDERS[0])
    v_mode = tk.StringVar(value="wzor")
    v_pat = tk.StringVar(value="{nazwa}")
    v_find, v_repl = tk.StringVar(), tk.StringVar()
    v_start = tk.StringVar(value="1")
    v_xlsx = tk.StringVar()

    f = ttk.Frame(root)
    f.pack(fill="x", **pad)
    ttk.Label(f, text="Folder:").grid(row=0, column=0, sticky="w", **pad)
    ttk.Entry(f, textvariable=v_dir, width=70).grid(row=0, column=1, columnspan=3, sticky="w", **pad)
    ttk.Button(f, text="Folder...", command=lambda: v_dir.set(
        filedialog.askdirectory() or v_dir.get())).grid(row=0, column=4, **pad)
    ttk.Label(f, text="Tylko pliki:").grid(row=1, column=0, sticky="w", **pad)
    ttk.Entry(f, textvariable=v_mask, width=12).grid(row=1, column=1, sticky="w", **pad)
    ttk.Label(f, text="np. *.pdf   kolejnosc:").grid(row=1, column=2, sticky="e", **pad)
    ttk.Combobox(f, textvariable=v_order, values=ORDERS, state="readonly",
                 width=16).grid(row=1, column=3, sticky="w", **pad)

    g = ttk.LabelFrame(root, text="Wedlug wzoru")
    g.pack(fill="x", **pad)
    ttk.Radiobutton(g, text="", variable=v_mode, value="wzor").grid(row=0, column=0)
    ttk.Label(g, text="Nowa nazwa:").grid(row=0, column=1, sticky="w", **pad)
    ttk.Entry(g, textvariable=v_pat, width=40).grid(row=0, column=2, sticky="w", **pad)
    ttk.Label(g, text="Nr od:").grid(row=0, column=3, sticky="e", **pad)
    ttk.Entry(g, textvariable=v_start, width=6).grid(row=0, column=4, sticky="w", **pad)
    ttk.Label(g, text="pola: {nazwa} {nr} {nr:3} = 001 {data} {folder}  (rozszerzenie zostaje)",
              foreground="gray").grid(row=1, column=2, columnspan=3, sticky="w", padx=6)
    ttk.Label(g, text="Zamien w nazwie:").grid(row=2, column=1, sticky="w", **pad)
    ttk.Entry(g, textvariable=v_find, width=40).grid(row=2, column=2, sticky="w", **pad)
    ttk.Label(g, text="na:").grid(row=2, column=3, sticky="e", **pad)
    ttk.Entry(g, textvariable=v_repl, width=20).grid(row=2, column=4, sticky="w", **pad)

    h = ttk.LabelFrame(root, text="Wedlug listy w Excelu (kolumna A = obecna nazwa, B = nowa)")
    h.pack(fill="x", **pad)
    ttk.Radiobutton(h, text="", variable=v_mode, value="excel").grid(row=0, column=0)
    ttk.Entry(h, textvariable=v_xlsx, width=62).grid(row=0, column=1, sticky="w", **pad)
    ttk.Button(h, text="Wybierz...", command=lambda: (v_xlsx.set(filedialog.askopenfilename(
        filetypes=[("Excel", "*.xlsx")]) or v_xlsx.get()), v_mode.set("excel"))
    ).grid(row=0, column=2, **pad)

    def make_list():
        folder = v_dir.get().strip('" ')
        if not os.path.isdir(folder):
            return status("Wskaz istniejacy folder.")
        path = filedialog.asksaveasfilename(defaultextension=".xlsx",
                                            initialfile="lista_plikow.xlsx",
                                            filetypes=[("Excel", "*.xlsx")])
        if path:
            names = list_files(folder, v_mask.get(), v_order.get())
            write_list(folder, names, path)
            v_xlsx.set(os.path.normpath(path))
            v_mode.set("excel")
            status("Zapisano liste %d plikow. Wpisz nowe nazwy w kolumnie B, zapisz plik"
                   " i kliknij Podglad." % len(names))
            os.startfile(path)
    ttk.Button(h, text="Utworz liste plikow...", command=make_list).grid(row=0, column=3, **pad)

    tree = ttk.Treeview(root, columns=("old", "new", "note"), show="headings", height=14)
    for col, text, w in (("old", "Obecna nazwa", 330), ("new", "Nowa nazwa", 330),
                         ("note", "Uwagi", 200)):
        tree.heading(col, text=text)
        tree.column(col, width=w)
    tree.tag_configure("err", background="#FFC7CE")
    tree.tag_configure("same", foreground="gray")
    tree.pack(fill="both", expand=True, **pad)

    msg = ttk.Label(root, text="")
    msg.pack(fill="x", padx=8)

    def status(text):
        msg.config(text=text)

    def current_plan():
        folder = v_dir.get().strip('" ')
        if not os.path.isdir(folder):
            raise ValueError("Wskaz istniejacy folder.")
        if v_mode.get() == "excel":
            if not os.path.isfile(v_xlsx.get().strip('" ')):
                raise ValueError("Wskaz plik Excela z lista nazw.")
            return folder, plan_excel(v_xlsx.get().strip('" '))
        try:
            start = int(v_start.get())
        except ValueError:
            raise ValueError("'Nr od' musi byc liczba.")
        names = list_files(folder, v_mask.get(), v_order.get())
        return folder, plan_pattern(folder, names, v_pat.get(), v_find.get(), v_repl.get(), start)

    def preview():
        tree.delete(*tree.get_children())
        try:
            folder, plan = current_plan()
        except Exception as e:
            status("BLAD: %s" % e)
            return None
        checked = check(folder, plan)
        for old, new, note in checked:
            tag = "same" if note == "bez zmian" else "err" if note else ""
            tree.insert("", "end", values=(old, new, note), tags=(tag,))
        n_err = len(errors(checked))
        n_ok = sum(1 for c in checked if c[2] == "")
        status("Do zmiany: %d, bez zmian: %d, bledy: %d%s" % (
            n_ok, len(checked) - n_ok - n_err, n_err,
            " - popraw bledy (na czerwono), zmiana nazw jest zablokowana" if n_err else ""))
        return folder, plan, n_ok, n_err

    def do_apply():
        res = preview()
        if not res:
            return
        folder, plan, n_ok, n_err = res
        if n_err or not n_ok:
            return
        if not messagebox.askyesno("Zmiana nazw", "Zmienic nazwy %d plikow?\n"
                                   "Mozna to potem cofnac przyciskiem 'Cofnij'." % n_ok):
            return
        try:
            n = apply(folder, plan)
            preview()
            status("Zmieniono nazwy %d plikow." % n)
        except Exception as e:
            status("BLAD: %s - nic nie zostalo zmienione." % e)

    def do_undo():
        last = last_change()
        if last is None:
            return status("Brak zmian do cofniecia.")
        if not messagebox.askyesno("Cofnij", "Przywrocic poprzednie nazwy %d plikow w:\n%s ?"
                                   % (len(last[2]), last[1])):
            return
        try:
            folder, n = undo_last()
        except Exception as e:
            return status("BLAD: %s" % e)
        status("Cofnieto zmiane nazw %d plikow w: %s" % (n, folder))
        preview()

    b = ttk.Frame(root)
    b.pack(pady=6)
    ttk.Button(b, text="Podglad", command=preview).pack(side="left", padx=6)
    ttk.Button(b, text="Zmien nazwy", command=do_apply).pack(side="left", padx=6)
    ttk.Button(b, text="Cofnij ostatnia zmiane", command=do_undo).pack(side="left", padx=6)

    if "--selftest" in sys.argv:
        root.after(200, root.destroy)
    import aktualizacja
    aktualizacja.start(root, "DawidBochno/Zmiana-nazw", "main", "zmiana_nazw.py")
    root.mainloop()


# --------------------------------------------------------------- selftest ----


def selftest():
    import shutil
    import tempfile

    assert sorted(["skan10.pdf", "skan2.pdf", "Skan1.pdf"], key=natural_key) == \
        ["Skan1.pdf", "skan2.pdf", "skan10.pdf"]
    assert name_error("ok.pdf") is None and name_error("a:b.pdf") and name_error("CON.txt")
    assert name_error("koniec. ") and name_error("")

    tmp = tempfile.mkdtemp()
    d = os.path.join(tmp, "Sprawa 12")
    os.makedirs(d)
    for n in ("skan10.pdf", "skan2.pdf", "skan1.pdf", "notatka.txt"):
        open(os.path.join(d, n), "w").close()
    os.makedirs(os.path.join(d, "podfolder.pdf"))
    names = list_files(d, "*.PDF")
    assert names == ["skan1.pdf", "skan2.pdf", "skan10.pdf"], names

    plan = plan_pattern(d, names, "{folder} - {nr:3} {nazwa}", "skan", "str", 1)
    assert plan[2] == ("skan10.pdf", "Sprawa 12 - 003 str10.pdf"), plan
    today = datetime.date.today().isoformat()
    assert plan_pattern(d, names[:1], "{data}_{nazwa}")[0][1] in (
        today + "_skan1.pdf", datetime.date.fromtimestamp(
            os.path.getmtime(os.path.join(d, "skan1.pdf"))).isoformat() + "_skan1.pdf")
    try:
        plan_pattern(d, names, "{autor}")
        raise AssertionError("nieznane pole przeszlo")
    except ValueError:
        pass

    # konflikty
    c = check(d, [("skan1.pdf", "notatka.txt"), ("skan2.pdf", "x.pdf"), ("skan10.pdf", "X.pdf"),
                  ("brak.pdf", "y.pdf"), ("notatka.txt", "notatka.txt")])
    notes = [x[2] for x in c]
    assert notes == ["dwa pliki dostałyby tę samą nazwę", "dwa pliki dostałyby tę samą nazwę", "dwa pliki dostałyby tę samą nazwę",
                     "brak takiego pliku w folderze", "bez zmian"], notes
    assert check(d, [("skan1.pdf", "notatka.txt")])[0][2] == "plik o tej nazwie już istnieje"
    hist = os.path.join(tmp, "hist")
    try:
        apply(d, [("skan1.pdf", "notatka.txt")], hist)
        raise AssertionError("apply zignorowal konflikt")
    except ValueError:
        pass

    # zamiana miejscami + zmiana wielkosci liter + cofniecie
    with open(os.path.join(d, "skan1.pdf"), "w") as fh:
        fh.write("jeden")
    n = apply(d, [("skan1.pdf", "skan2.pdf"), ("skan2.pdf", "skan1.pdf"),
                  ("notatka.txt", "Notatka.txt")], hist)
    assert n == 3
    with open(os.path.join(d, "skan2.pdf")) as fh:
        assert fh.read() == "jeden"
    assert "Notatka.txt" in os.listdir(d)
    assert undo_last(hist) == (os.path.abspath(d), 3)
    with open(os.path.join(d, "skan1.pdf")) as fh:
        assert fh.read() == "jeden"
    assert "notatka.txt" in os.listdir(d) and undo_last(hist) is None

    # blad w trakcie -> nic nie zostaje zmienione
    real = os.rename
    calls = []

    def flaky(a, b):
        calls.append(a)
        if len(calls) == 3:
            raise OSError("plik zablokowany")
        real(a, b)
    os.rename = flaky
    try:
        _rename_all(d, [("skan1.pdf", "a.pdf"), ("skan2.pdf", "b.pdf")])
        raise AssertionError("brak bledu")
    except OSError:
        pass
    finally:
        os.rename = real
    assert sorted(os.listdir(d)) == ["notatka.txt", "podfolder.pdf", "skan1.pdf",
                                     "skan10.pdf", "skan2.pdf"], os.listdir(d)

    # Excel: lista -> nowe nazwy -> plan
    import openpyxl
    xlsx = os.path.join(tmp, "lista.xlsx")
    write_list(d, list_files(d), xlsx)
    wb = openpyxl.load_workbook(xlsx)
    ws = wb.active
    assert ws["A2"].value == "notatka.txt" and ws["B1"].value == "Nowa nazwa"
    ws["B2"] = "Notatka sluzbowa"
    ws["B4"] = "Strona 1.pdf"
    wb.save(xlsx)
    plan = plan_excel(xlsx)
    assert plan == [("notatka.txt", "Notatka sluzbowa.txt"), ("skan2.pdf", "Strona 1.pdf")], plan
    assert apply(d, plan, hist) == 2
    assert "Notatka sluzbowa.txt" in os.listdir(d)
    shutil.rmtree(tmp, ignore_errors=True)

    import aktualizacja
    aktualizacja.selftest()
    print("selftest OK")


if __name__ == "__main__":
    if "--selftest" in sys.argv:
        selftest()
        if "--gui" in sys.argv:
            gui()
    else:
        gui()
